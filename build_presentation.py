"""
SonicSentinel AI — Professional PowerPoint Presentation Deck Generator
Generates SonicSentinel_Project_Presentation.pptx with 20+ slides.
"""
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "documentation"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Theme
DARK_BG = RGBColor(7, 11, 18)
CYAN = RGBColor(0, 240, 255)
BLUE = RGBColor(0, 112, 243)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(138, 153, 173)

def add_blank_slide():
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG
    return slide

def add_header(slide, title, category="SONICSENTINEL AI"):
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p_cat = tf.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = CYAN
    
    p_title = tf.add_paragraph()
    p_title.text = title
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE

def add_card(slide, left, top, width, height, title, body_text, accent=CYAN):
    shape = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(13, 21, 39)
    shape.line.color.rgb = accent
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(14)
    p_title.font.bold = True
    p_title.font.color.rgb = accent
    
    p_body = tf.add_paragraph()
    p_body.text = body_text
    p_body.font.size = Pt(11)
    p_body.font.color.rgb = WHITE

print("Generating 20-Slide Presentation Deck...")

# Slide 1: Cover
s1 = add_blank_slide()
tx = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
tf = tx.text_frame
p1 = tf.paragraphs[0]
p1.text = "SonicSentinel AI"
p1.font.size = Pt(44)
p1.font.bold = True
p1.font.color.rgb = CYAN

p2 = tf.add_paragraph()
p2.text = "AcousticX Intelligence — Sound Event Classification & Critical Security System"
p2.font.size = Pt(20)
p2.font.color.rgb = WHITE

p3 = tf.add_paragraph()
p3.text = "\nFinal Year Project Presentation | Aptech TechWiz 7 — NextWave AI/ML Category"
p3.font.size = Pt(14)
p3.font.color.rgb = MUTED

# Slide 2: Problem Statement
s2 = add_blank_slide()
add_header(s2, "The Problem Statement: Acoustic Monitoring Gaps")
add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), "Surveillance Blind Spots", "Visual cameras require direct line of sight and adequate lighting. Critical incidents like gunshots, screams, glass breaking, or machinery failures often happen outside visual range.", CRIMSON:=RGBColor(255, 51, 102))
add_card(s2, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), "Manual Operator Fatigue", "Human guards cannot continuously listen to hours of multi-channel audio recordings. Incidents are frequently missed or detected too late for rapid emergency response.", AMBER:=RGBColor(255, 184, 0))

# Slide 3: Proposed Solution
s3 = add_blank_slide()
add_header(s3, "The Proposed Solution: SonicSentinel AI")
add_card(s3, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), "24/7 Listening Engine", "Replaces manual audio listening with automated 3.5s microphone chunk analysis and file upload processing.", CYAN)
add_card(s3, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8), "3-Model Python Ensemble", "Combines Random Forest (93.7%), SVM (91.6%), and 2D CNN (91.0%) for robust multi-model consensus decision making.", BLUE)
add_card(s3, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8), "Real-Time Critical Alerting", "Automated severity assessment (Critical, High, Medium, Low) with live dashboard alerts and manual review routing.", RGBColor(47, 224, 168))

# Slide 4: 10 Sound Categories
s4 = add_blank_slide()
add_header(s4, "10 Mandatory Sound Event Categories")
add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), "Security & Human Safety", "1. Gunshot (Critical)\n2. Panic Scream (Critical)\n3. Person Asking for Help (Critical)\n4. Aggression / Violent Conflict (High)\n5. Glass Breaking (High)\n6. Alarm or Siren (High)", CYAN)
add_card(s4, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), "Industrial & Environmental", "7. Machinery Fault (Medium)\n8. Vehicle Horn (Low)\n9. Animal Sound (Low)\n10. Background Noise (Low)", BLUE)

# Slide 5: System Architecture
s5 = add_blank_slide()
add_header(s5, "High-Level System Architecture")
add_card(s5, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "Unified Container Architecture", "• React Vite Frontend: User & Admin Dashboards with Waveform & Live Mic Visualizer\n• FastAPI Backend Gateway: REST Endpoints & Web Audio chunk ingestion\n• 3-Model Python ML Engine: 242 Feature Extractor + Scikit-Learn / Keras Evaluation\n• SQLite / MongoDB Storage: Hashed credentials, audit logs, and detection history", CYAN)

# Slide 6: Audio Preprocessing Pipeline
s6 = add_blank_slide()
add_header(s6, "Audio Preprocessing Pipeline")
add_card(s6, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "Standardized Signal Preparation", "1. Format Sniffing & Validation: Supports WAV, MP3, OGG, FLAC, AAC, WEBM\n2. Resampling & Mono Downmixing: Standardized to 22,050 Hz Mono\n3. Silence Trimming: librosa.effects.trim(top_db=35)\n4. Length Standardization: Padded or cropped to exactly 5.0s (110,250 samples)\n5. Peak Normalization: y = y / max(|y|) for invariant amplitude scaling", BLUE)

# Slide 7: Feature Extraction (242 Features)
s7 = add_blank_slide()
add_header(s7, "Feature Extraction (242 Acoustic Features)")
add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), "Tabular Features (Indices 0:242)", "• 80 MFCC Features: 40 Means + 40 Stds (n_mfcc=40)\n• 128 Mel64 Features: 64 Means + 64 Stds (n_mels=64)\n• 24 Chroma STFT: 12 Means + 12 Stds\n• 10 Spectral Features: ZCR, RMS, Centroid, Bandwidth, Rolloff", CYAN)
add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), "CNN Spectrogram Input", "• 128x216 Log Mel Spectrogram\n• n_fft=1024, hop_length=512, librosa.power_to_db\n• Min-Max Scaled to [0, 1] for 2D ConvNet evaluation", PURPLE:=RGBColor(123, 97, 255))

# Slide 8: Model 1 — Random Forest
s8 = add_blank_slide()
add_header(s8, "Model 1: Random Forest Classifier")
add_card(s8, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "Random Forest Architecture & Performance", "• Architecture: 400 Decision Trees with balanced class weights\n• Input: 242-dimensional acoustic feature vector\n• Accuracy: 93.67% (Top Performing Model)\n• Macro Precision: 94.50% | Macro Recall: 92.25% | Macro F1: 93.09%\n• Strength: High resistance to feature noise and class imbalance", CYAN)

# Slide 9: Model 2 — SVM Pipeline
s9 = add_blank_slide()
add_header(s9, "Model 2: Support Vector Machine (SVM)")
add_card(s9, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "SVM Pipeline Architecture & Performance", "• Architecture: StandardScaler + SVC (RBF Kernel, C=10, gamma='scale')\n• Probability Output: Enabled via Platt scaling for soft voting ensemble\n• Accuracy: 91.63%\n• Macro Precision: 92.19% | Macro Recall: 90.46% | Macro F1: 91.13%\n• Strength: Optimal decision boundary in high-dimensional feature space", PURPLE)

# Slide 10: Model 3 — 2D Convolutional Neural Network
s10 = add_blank_slide()
add_header(s10, "Model 3: 2D Convolutional Neural Network (CNN)")
add_card(s10, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "2D CNN Architecture & Performance", "• Architecture: 3 Conv2D Blocks (32, 64, 128 filters) + Batch Normalization + MaxPool2D\n• Regularization: Global Average Pooling + Dropout (0.25, 0.35)\n• Accuracy: 90.99%\n• Macro Precision: 90.90% | Macro Recall: 90.36% | Macro F1: 90.40%\n• Strength: Learns spatial-frequency patterns directly from Mel Spectrograms", RGBColor(47, 224, 168))

# Slide 11: 3-Model Ensemble & Decision Engine
s11 = add_blank_slide()
add_header(s11, "Multi-Model Ensemble & Decision Engine")
add_card(s11, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "Weighted Probability Soft Voting", "• Weighted Formula: P_ensemble = 0.40 * P_RF + 0.35 * P_SVM + 0.25 * P_CNN\n• Model Agreement Status:\n  - Acceptable Match: All 3 models agree on top predicted class\n  - Weak Match: 2 out of 3 models agree\n  - Model Disagreement: All 3 models differ (triggers Manual Review Queue)", CYAN)

# Slide 12: Critical Alert & Severity System
s12 = add_blank_slide()
add_header(s12, "Critical Alert & Severity System")
add_card(s12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), "Severity Rules", "• Critical: Gunshot, Panic Scream, Person Asking for Help (Conf >= 70%)\n• High: Glass Breaking, Alarm/Siren, Aggression\n• Medium: Machinery Fault\n• Low: Vehicle Horn, Animal Sound, Background Noise", CRIMSON)
add_card(s12, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), "Alert Cooldown & Deduplication", "Prevents false alarm fatigue by enforcing a 60-second cooldown deduplication window for repeated high-severity events in the same location stream.", AMBER)

# Slide 13: Live Microphone Monitoring
s13 = add_blank_slide()
add_header(s13, "Live Microphone Monitoring & Real-Time Stream")
add_card(s13, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "Web Audio & MediaRecorder Stream Processing", "• Real-Time Audio Canvas: AudioContext + AnalyserNode (2048 FFT) for real-time waveform visualizer\n• Continuous Mic Slicing: MediaRecorder packages 3.5s Opus/WEBM chunks\n• Real-Time Post: Chunks posted to /live/analyze without interrupting mic stream\n• Instant Toast Alerts: High/Critical detections trigger instant audio-visual alerts", CYAN)

# Slide 14: Authentication & Security Architecture
s14 = add_blank_slide()
add_header(s14, "Security, Auth & Role-Based Access Control")
add_card(s14, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), "Public Signup RBAC Hardening", "• Public signup ALWAYS assigns role = 'user'\n• Client attempts to send role='admin' or isAdmin=true are strictly ignored\n• Dedicated CLI script (python -m app.create_admin) for admin provisioning", CYAN)
add_card(s14, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8), "Server-Side Authorization", "• require_admin FastAPI dependency guards all /api/admin/* endpoints\n• Unprivileged requests receive HTTP 403 Forbidden\n• Passwords hashed with PBKDF2-SHA256 (310,000 iterations)", BLUE)

# Slide 15: Admin Dashboard
s15 = add_blank_slide()
add_header(s15, "Admin Dashboard & Management Subsystem")
add_card(s15, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "System-Wide Overview & Control", "• Overview: Real-time platform stats, severity distribution charts, class breakdown\n• User Management: Search, filter, activate/deactivate, delete users with confirmation\n• Detections & Alerts: Searchable platform-wide audit log\n• System Health: Database status, model status, uptime metrics", PURPLE)

# Slide 16: Database Design (ERD)
s16 = add_blank_slide()
add_header(s16, "Database Architecture & Collections")
add_card(s16, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "MongoDB / SQLite Compatible Schema Layer", "• users: _id, name, email, password_hash, role (user/admin), active (bool), created_at\n• sessions: token_hash, user_id, expires_at\n• detections: user_id, audio_filename, classification, confidence, severity, models, audio_quality\n• alerts: user_id, detection_id, severity, label, message, read, resolved\n• reviews: user_id, detection_id, final_classification, notes, status", CYAN)

# Slide 17: Testing & Verification
s17 = add_blank_slide()
add_header(s17, "Testing & Automated Security Verification")
add_card(s17, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "12-Point Security & Functional Test Suite", "• 100% Pass Rate across 12 automated test cases (test_security.py)\n• Verified: Public signup role isolation, admin bypass rejection, HTTP 403 enforcement, password hashing, and user data isolation\n• Production Frontend Build: Vite production build transformed 3,096 modules with 0 errors", RGBColor(47, 224, 168))

# Slide 18: Railway Cloud Deployment
s18 = add_blank_slide()
add_header(s18, "Railway Cloud Deployment Architecture")
add_card(s18, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "Single Container Unified Deployment", "• Multi-Stage Dockerfile: Stage 1 Node.js build + Stage 2 Python 3.11-slim server\n• System Audio Dependencies: ffmpeg and libsndfile1 pre-installed\n• Auto Port Binding: Binds dynamically to Railway $PORT\n• Static Asset Serving: FastAPI mounts compiled React dist/ assets for SPA routing", CYAN)

# Slide 19: SRS Compliance Matrix Summary
s19 = add_blank_slide()
add_header(s19, "SRS Compliance Matrix Summary")
add_card(s19, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "99 Mandatory Requirements Traceability", "• Functional Requirements (80): 77 Implemented, 3 Deferred (GTM optional adapter)\n• Non-Functional Requirements (5): 5 Implemented & Verified\n• Competition Integrity / Anti-Shortcut Requirements (14): 14 Implemented & Verified\n• Overall Compliance Rate: 96 out of 99 (97.0%)", BLUE)

# Slide 20: Conclusion & Q&A
s20 = add_blank_slide()
add_header(s20, "Conclusion & Project Summary")
add_card(s20, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.8), "SonicSentinel AI Ready for Production", "• Delivered 3-Model Python Ensemble (RF 93.7%, SVM 91.6%, 2D CNN 91.0%)\n• Implemented complete RBAC security and Admin Dashboard\n• Enabled real-time live microphone monitoring and critical event alerting\n• Configured 100% verified Railway Cloud deployment\n\nThank you! Questions & Discussion.", CYAN)

# Save Presentation
prs.save(DOCS / "SonicSentinel_Project_Presentation.pptx")
print("Saved SonicSentinel_Project_Presentation.pptx")
