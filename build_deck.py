from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
O=Path(r"E:\SonicSentinel AI\documentation")
S=[
("COVER","SonicSentinel AI","Project documentation","Evidence-led technical showcase"),
("EXECUTIVE SUMMARY","Audio intelligence with a reviewable path","Product","React/Vite frontend, FastAPI service, persistence, and ML boundary"),
("PROBLEM","Operational audio is difficult to triage at scale","Context","Detection needs confidence, severity, auditability, and human review"),
("OBJECTIVES","Detect, compare, alert, and document","Objectives","Upload/live capture, model comparison, alerts, review workflow"),
("PROPOSED SOLUTION","A dual-engine monitoring workflow","Solution","Input -> validation -> inference -> confidence -> alert/review"),
("SYSTEM ARCHITECTURE","Separated application boundaries","Verified source","Frontend -> FastAPI -> SQLite/uploads; model boundary is explicit"),
("TECHNOLOGY STACK","The inspected implementation surface","Stack","React, Vite, TypeScript, Framer Motion, Recharts, FastAPI, Python, SQLite"),
("REQUIREMENTS","SRS mapping is part of the deliverable","Evidence","See SRS_COMPLIANCE_MATRIX.md; statuses are mapped to inspected evidence"),
("DATASET","Runtime audio inventory","Observed","WAV upload inventory generated from ml-service/data/uploads"),
("DATASET DISTRIBUTION","Labels are not present in upload filenames","Status","Class distribution: PENDING VERIFICATION"),
("DATASET QUALITY","Readable WAV metadata was inventoried","Evidence","See dataset/audio-inventory.csv and quality-report.md"),
("AUDIO PREPROCESSING","Input validation precedes inference","Pipeline","Container validation and feature extraction implementation require runtime evidence"),
("FEATURE EXTRACTION","Feature provenance must be reproducible","Status","Exact training features and parameters: PENDING VERIFICATION"),
("RANDOM FOREST","Model evidence boundary","Status","RF metrics/artifact: RESULTS PENDING"),
("SVM","Model evidence boundary","Status","SVM metrics/artifact: RESULTS PENDING"),
("CNN","Model evidence boundary","Status","CNN metrics/artifact: RESULTS PENDING"),
("TRAINING PIPELINE","Dataset to deployable artifact","Status","Training run artifacts are not present in the repository"),
("MODEL EVALUATION","Evidence before numbers","Status","No accuracy, precision, recall, F1, or confusion matrix is invented"),
("CONFUSION MATRICES","Class-level diagnostics","Status","RF/SVM/CNN matrices: RESULTS PENDING"),
("MODEL COMPARISON","Agreement is a workflow decision","Status","Comparison CSV is explicitly marked RESULTS PENDING"),
("PREDICTION & CONFIDENCE","Confidence should route uncertainty","Implementation","Detections retain prediction data through API routes; saved samples pending"),
("CRITICAL ALERT SYSTEM","Severity becomes an operator action","Routes","Alerts exposed by GET /alerts and PATCH /alerts/{alert_id}"),
("APPLICATION SURFACES","Frontend route and screen inventory","UI","Landing, auth, dashboard, analyze, live, history, alerts, reports, models, settings"),
("DATABASE / ERD","Persistence boundary","Observed","SQLite files and entities are implemented in ml-service/app/sqlite_db.py"),
("API ARCHITECTURE","FastAPI route catalog","Verified","Full method/path list is in API_DOCUMENTATION.md"),
("AUTHENTICATION","Session-aware protected access","Routes","Register, login, me, profile, logout, reset routes are present"),
("SECURITY","Keep secrets and microphone consent server-bound","Status","Security test evidence: PENDING VERIFICATION"),
("TESTING","Build evidence exists; functional evidence remains","Verified","npm run build passed; test suite artifacts are pending"),
("SRS COMPLIANCE","Traceability from requirement to evidence","Matrix","See SRS_COMPLIANCE_MATRIX.md"),
("RESULTS","Separate observed facts from open claims","Summary","22 diagrams created; model metrics and screenshots pending"),
("LIMITATIONS","What this package does not claim","Honesty","No fabricated model numbers, screenshots, or runtime test outcomes"),
("FUTURE WORK","Close the evidence gaps","Next steps","Capture browser screens, run ML evaluation, add labelled dataset manifest"),
("CONCLUSION","A credible technical record","Close","Every open claim has an explicit evidence status and file location")
]
def tx(s,x,y,w,h,t,size,c,b=False):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));p=q.text_frame.paragraphs[0];r=p.add_run();r.text=t;r.font.name="Aptos";r.font.size=Pt(size);r.font.bold=b;r.font.color.rgb=RGBColor(*c)
P=Presentation();P.slide_width=Inches(13.333);P.slide_height=Inches(7.5)
for i,(k,t,u,bod) in enumerate(S,1):
 s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(8,14,28)
 z=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(.18),P.slide_height);z.fill.solid();z.fill.fore_color.rgb=RGBColor(58,205,222);z.line.fill.background()
 c=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.72),Inches(1.35),Inches(11.85),Inches(4.75));c.fill.solid();c.fill.fore_color.rgb=RGBColor(18,29,52);c.line.fill.background()
 tx(s,.85,.55,11.5,.3,k,11,(58,205,222),True);tx(s,.85,1.75,10.8,.8,t,27,(240,244,250),True);tx(s,.85,2.7,10.5,.45,u,14,(157,173,196));tx(s,.85,3.45,10.8,1.5,bod,20,(240,244,250));tx(s,.85,6.8,8,.25,"SONICSENTINEL AI | EVIDENCE REGISTER",9,(157,173,196),True);tx(s,11.8,6.8,.5,.25,f"{i:02d}",9,(58,205,222),True)
P.save(O/"SonicSentinel_Project_Documentation.pptx")
with PdfPages(O/"SonicSentinel_Project_Documentation.pdf") as pdf:
 for i,(k,t,u,bod) in enumerate(S,1):
  f=plt.figure(figsize=(13.333,7.5),facecolor="#080e1c");a=f.add_axes([0,0,1,1]);a.axis("off");a.add_patch(plt.Rectangle((0,0),.018,1,color="#3acdde",transform=a.transAxes));a.add_patch(plt.Rectangle((.055,.18),.89,.63,color="#121d34",transform=a.transAxes));a.text(.065,.91,k,color="#3acdde",fontsize=11,fontweight="bold",transform=a.transAxes);a.text(.065,.72,t,color="#f0f4fa",fontsize=25,fontweight="bold",transform=a.transAxes);a.text(.065,.62,u,color="#9dadc4",fontsize=14,transform=a.transAxes);a.text(.065,.47,bod,color="#f0f4fa",fontsize=19,wrap=True,transform=a.transAxes);a.text(.065,.07,"SONICSENTINEL AI | EVIDENCE REGISTER",color="#9dadc4",fontsize=9,fontweight="bold",transform=a.transAxes);a.text(.88,.07,f"{i:02d}",color="#3acdde",fontsize=9,fontweight="bold",transform=a.transAxes);pdf.savefig(f);plt.close(f)
print("slides",len(S))
